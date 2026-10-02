"""
Modelos de recolha de dados:
- excel_template: livro Excel formatado, com validação de dados, pronto a preencher e a carregar na app.
- questionnaire_docx: questionário em Word para cada decisor preencher (em papel ou digitalmente).
"""

from __future__ import annotations

import io

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PINE = "2E5E4E"
MIST = "EEF2F0"
AMBER_SOFT = "F6E7C8"
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=PINE)
INPUT_FILL = PatternFill("solid", fgColor="FFFFFF")
HELP_FONT = Font(italic=True, color="4B5B56", size=9)
THIN = Border(*(Side(style="thin", color="C9D6D0"),) * 4)

SCALE = [(1, "Igual importância"), (2, "Entre igual e moderada"), (3, "Moderadamente mais importante"),
         (4, "Entre moderada e forte"), (5, "Fortemente mais importante"), (6, "Entre forte e muito forte"),
         (7, "Muito fortemente mais importante"), (8, "Entre muito forte e extrema"),
         (9, "Extremamente mais importante")]

EXTRA_ROWS = 40  # linhas em branco com validação para acrescentar dados


def _header(ws, row, values, widths=None):
    for j, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=j, value=v)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, THIN
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _cells(ws, r1, r2, c1, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = THIN


def excel_template(crit_df: pd.DataFrame, dms_df: pd.DataFrame, bo: pd.DataFrame | None = None,
                   ow: pd.DataFrame | None = None, alt: pd.DataFrame | None = None,
                   meta: dict | None = None) -> bytes:
    meta = meta or {}
    crit_df = crit_df.copy()
    for col in ("Unidade", "Descrição"):
        if col not in crit_df:
            crit_df[col] = ""
    crit_df["Critério"] = crit_df["Critério"].astype(str).str.strip()
    crit_df = crit_df[(crit_df["Critério"] != "") & (crit_df["Critério"] != "nan")]
    crit_df = crit_df[["Critério", "Tipo", "Unidade", "Descrição"]].fillna("").reset_index(drop=True)
    dms_df = dms_df.copy()
    dms_df["Decisor"] = dms_df["Decisor"].astype(str).str.strip()
    dms_df = dms_df[(dms_df["Decisor"] != "") & (dms_df["Decisor"] != "nan")].reset_index(drop=True)
    criteria = crit_df["Critério"].astype(str).tolist()
    types = dict(zip(crit_df["Critério"], crit_df["Tipo"]))
    units = dict(zip(crit_df["Critério"], crit_df["Unidade"].fillna("")))
    dms = dms_df["Decisor"].astype(str).tolist()
    n = len(criteria)

    wb = Workbook()

    # --- Instruções ---------------------------------------------------------
    ws = wb.active
    ws.title = "Instrucoes"
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 110
    steps = [
        ("Modelo de recolha de dados — Bayesian Best-Worst Method", True),
        ("Preencha as folhas pela ordem abaixo e carregue este ficheiro na app (Origem dos dados → Carregar ficheiro).", False),
        ("", False),
        ("1. Projeto — título da decisão, setor/organização e contexto. Aparecem no relatório e orientam a interpretação.", False),
        ("2. Criterios — um critério por linha. Tipo: Benefício (mais é melhor) ou Custo (menos é melhor). "
         "Indique a unidade (€, dias, %, escala 1-10).", False),
        ("3. Decisores — uma pessoa por linha, com a função ou área.", False),
        ("4. Comparacoes — duas linhas por decisor: Tipo = BO e Tipo = OW. Os nomes das colunas têm de ser IGUAIS "
         "aos nomes da folha Criterios.", False),
        ("     BO (Best-to-Others): quanto o MELHOR critério do decisor é mais importante do que cada critério. "
         "O melhor leva 1.", False),
        ("     OW (Others-to-Worst): quanto cada critério é mais importante do que o PIOR critério do decisor. "
         "O pior leva 1.", False),
        ("     Use valores inteiros de 1 a 9 (ver folha Escala). As células com 1 ficam destacadas.", False),
        ("     O valor do melhor face ao pior deve ser igual nas duas linhas (BO na coluna do pior = OW na coluna do melhor).", False),
        ("5. Alternativas (opcional) — opções a ordenar, com o desempenho em cada critério nas unidades indicadas.", False),
        ("", False),
        ("Para acrescentar um critério: nova linha em Criterios + nova coluna com o mesmo nome em Comparacoes e Alternativas.", False),
        ("Para acrescentar um decisor: nova linha em Decisores + duas linhas (BO e OW) em Comparacoes.", False),
        ("A coluna Notas é ignorada no cálculo e pode ser usada para comentários.", False),
        ("Dica: use o questionário Word (botão na app) para recolher as respostas de cada decisor e transcreva-as aqui.", False),
    ]
    for i, (t, bold) in enumerate(steps, start=1):
        c = ws.cell(row=i, column=2, value=t)
        c.font = Font(bold=bold, size=14 if bold else 11, color=PINE if bold else "1B2A26")
        c.alignment = Alignment(wrap_text=True)

    # --- Projeto ------------------------------------------------------------
    ws = wb.create_sheet("Projeto")
    _header(ws, 1, ["Campo", "Valor", "Ajuda"], [24, 70, 55])
    rows = [("Título da decisão", meta.get("title", ""), "Ex.: Seleção de seguradora para frota automóvel"),
            ("Setor / organização", meta.get("sector", ""), "Ex.: Seguros — gestão de sinistros de acidentes"),
            ("Contexto", meta.get("context", ""), "Problema, objetivo, restrições e quem decide.")]
    for i, (k, v, h) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        cv = ws.cell(row=i, column=2, value=v)
        cv.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=3, value=h).font = HELP_FONT
    ws.row_dimensions[4].height = 90
    _cells(ws, 2, 4, 1, 3)

    # --- Critérios ----------------------------------------------------------
    ws = wb.create_sheet("Criterios")
    _header(ws, 1, ["Critério", "Tipo", "Unidade", "Descrição"], [32, 14, 22, 70])
    for i, r in enumerate(crit_df.itertuples(index=False), start=2):
        ws.cell(row=i, column=1, value=r[0])
        ws.cell(row=i, column=2, value=r[1])
        ws.cell(row=i, column=3, value=r[2] if isinstance(r[2], str) else "")
        ws.cell(row=i, column=4, value=r[3] if isinstance(r[3], str) else "")
    last = 1 + n + EXTRA_ROWS
    dv = DataValidation(type="list", formula1='"Benefício,Custo"', allow_blank=True,
                        promptTitle="Tipo de critério", prompt="Benefício: mais é melhor. Custo: menos é melhor.",
                        showErrorMessage=True, errorTitle="Valor inválido", error="Escolha Benefício ou Custo.")
    ws.add_data_validation(dv)
    dv.add(f"B2:B{last}")
    _cells(ws, 2, last, 1, 4)

    # --- Decisores ----------------------------------------------------------
    ws = wb.create_sheet("Decisores")
    _header(ws, 1, ["Decisor", "Função"], [36, 36])
    for i, r in enumerate(dms_df.itertuples(index=False), start=2):
        ws.cell(row=i, column=1, value=r[0])
        ws.cell(row=i, column=2, value=r[1] if len(r) > 1 and isinstance(r[1], str) else "")
    _cells(ws, 2, 1 + len(dms) + EXTRA_ROWS, 1, 2)

    # --- Comparações --------------------------------------------------------
    ws = wb.create_sheet("Comparacoes")
    head = ["Decisor", "Tipo"] + criteria + ["Notas"]
    _header(ws, 1, head, [30, 8] + [max(12, min(24, len(c) + 4)) for c in criteria] + [40])
    for j, c in enumerate(criteria, start=3):
        ws.cell(row=1, column=j).comment = Comment(
            f"{c} ({types.get(c, '')})\nBO: melhor vs. {c}\nOW: {c} vs. pior", "Modelo BWM")
    r = 2
    for dm in dms:
        for kind, mat in (("BO", bo), ("OW", ow)):
            ws.cell(row=r, column=1, value=dm)
            ws.cell(row=r, column=2, value=kind)
            if mat is not None and dm in mat.index:
                for j, c in enumerate(criteria, start=3):
                    val = mat.loc[dm, c] if c in mat.columns else None
                    val = pd.to_numeric(val, errors="coerce")
                    ws.cell(row=r, column=j, value=None if pd.isna(val) else int(val))
            r += 1
    last = r - 1 + EXTRA_ROWS
    c1, c2 = get_column_letter(3), get_column_letter(2 + n)
    dv_t = DataValidation(type="list", formula1='"BO,OW"', allow_blank=True, showErrorMessage=True,
                          error="Use BO ou OW.", promptTitle="Tipo de linha",
                          prompt="BO = Best-to-Others; OW = Others-to-Worst")
    dv_v = DataValidation(type="whole", operator="between", formula1="1", formula2="9", allow_blank=True,
                          showErrorMessage=True, errorTitle="Fora da escala", error="Valor inteiro entre 1 e 9.",
                          promptTitle="Escala 1-9", prompt="1 = igual importância … 9 = extremamente mais importante")
    ws.add_data_validation(dv_t)
    ws.add_data_validation(dv_v)
    dv_t.add(f"B2:B{last}")
    dv_v.add(f"{c1}2:{c2}{last}")
    ws.conditional_formatting.add(f"{c1}2:{c2}{last}",
                                  CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor=AMBER_SOFT),
                                             font=Font(bold=True)))
    for rr in range(2, last + 1):
        if rr % 2 == 0:
            for cc in range(1, 3 + n):
                ws.cell(row=rr, column=cc).fill = PatternFill("solid", fgColor=MIST)
    _cells(ws, 2, last, 1, 3 + n)

    # --- Alternativas -------------------------------------------------------
    ws = wb.create_sheet("Alternativas")
    head = ["Alternativa"] + criteria + ["Notas"]
    _header(ws, 1, head, [32] + [max(14, min(26, len(c) + 6)) for c in criteria] + [40])
    for j, c in enumerate(criteria, start=2):
        arrow = "↑ mais é melhor" if types.get(c) == "Benefício" else "↓ menos é melhor"
        u = f" — unidade: {units.get(c)}" if units.get(c) else ""
        ws.cell(row=1, column=j).comment = Comment(f"{c}: {arrow}{u}", "Modelo BWM")
    ws.insert_rows(2)
    ws.cell(row=2, column=1, value="(unidade / sentido)").font = HELP_FONT
    for j, c in enumerate(criteria, start=2):
        arrow = "↑" if types.get(c) == "Benefício" else "↓"
        ws.cell(row=2, column=j, value=f"{arrow} {units.get(c) or ''}".strip()).font = HELP_FONT
    ws.freeze_panes = "B3"
    r = 3
    if alt is not None and len(alt):
        for name, row in alt.iterrows():
            ws.cell(row=r, column=1, value=str(name))
            for j, c in enumerate(criteria, start=2):
                v = pd.to_numeric(row.get(c), errors="coerce")
                ws.cell(row=r, column=j, value=None if pd.isna(v) else float(v))
            r += 1
    last = max(r - 1, 2) + EXTRA_ROWS
    dv_n = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="-1000000000", allow_blank=True,
                          showErrorMessage=True, error="Introduza um número.")
    ws.add_data_validation(dv_n)
    dv_n.add(f"B3:{get_column_letter(1 + n)}{last}")
    _cells(ws, 3, last, 1, 2 + n)

    # --- Escala -------------------------------------------------------------
    ws = wb.create_sheet("Escala")
    _header(ws, 1, ["Valor", "Significado"], [8, 45])
    for i, (v, t) in enumerate(SCALE, start=2):
        ws.cell(row=i, column=1, value=v)
        ws.cell(row=i, column=2, value=t)
    _cells(ws, 2, 10, 1, 2)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Questionário Word para os decisores
# ---------------------------------------------------------------------------


def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = h
        _shade(c, PINE)
        for p in c.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(9.5)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def questionnaire_docx(crit_df: pd.DataFrame, meta: dict | None = None) -> bytes:
    meta = meta or {}
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(1.8)

    doc.add_heading("Questionário de avaliação de critérios", 0)
    doc.add_paragraph("Método: Best-Worst Method (BWM)").runs[0].font.color.rgb = RGBColor(0x4B, 0x5B, 0x56)
    if meta.get("title"):
        p = doc.add_paragraph()
        p.add_run("Decisão: ").bold = True
        p.add_run(meta["title"])
    if meta.get("sector"):
        p = doc.add_paragraph()
        p.add_run("Setor / organização: ").bold = True
        p.add_run(meta["sector"])
    if meta.get("context"):
        doc.add_paragraph(meta["context"])

    doc.add_heading("1. Identificação", 1)
    _table(doc, ["Campo", "Resposta"], [["Nome", ""], ["Função / área", ""], ["Data", ""]], [5, 12])

    doc.add_heading("2. Critérios em avaliação", 1)
    doc.add_paragraph("Leia a descrição de cada critério antes de responder. Avalie a importância de cada critério "
                      "para esta decisão, não o desempenho de uma alternativa concreta.")
    rows = [[r["Critério"], r.get("Descrição", "") or ""] for _, r in crit_df.iterrows()]
    _table(doc, ["Critério", "Descrição"], rows, [5.5, 11.5])

    doc.add_heading("3. Escala de comparação", 1)
    _table(doc, ["Valor", "Significado"], [[v, t] for v, t in SCALE], [2, 10])

    crit = crit_df["Critério"].astype(str).tolist()
    doc.add_heading("4. Melhor e pior critério", 1)
    doc.add_paragraph("Assinale com X o critério MAIS importante (melhor) e o MENOS importante (pior). "
                      "Escolha apenas um em cada coluna; não podem ser o mesmo.")
    _table(doc, ["Critério", "Melhor (X)", "Pior (X)"], [[c, "", ""] for c in crit], [9, 4, 4])

    doc.add_heading("5. Comparação Best-to-Others (BO)", 1)
    doc.add_paragraph("Pergunta: «Quanto é o MEU MELHOR critério mais importante do que cada um destes?» "
                      "Responda de 1 a 9. Na linha do próprio melhor critério escreva 1.")
    _table(doc, ["Critério", "Valor (1-9)"], [[c, ""] for c in crit], [11, 5])

    doc.add_heading("6. Comparação Others-to-Worst (OW)", 1)
    doc.add_paragraph("Pergunta: «Quanto é cada um destes critérios mais importante do que o MEU PIOR critério?» "
                      "Responda de 1 a 9. Na linha do próprio pior critério escreva 1.")
    _table(doc, ["Critério", "Valor (1-9)"], [[c, ""] for c in crit], [11, 5])

    doc.add_heading("7. Verificação rápida antes de entregar", 1)
    for t in [
        "Na tabela BO, o melhor critério tem 1 e o pior tem o valor mais alto.",
        "Na tabela OW, o pior critério tem 1 e o melhor tem o valor mais alto.",
        "O valor «melhor vs. pior» é igual nas duas tabelas.",
        "Coerência: se o melhor é 2× mais importante que A e A é 3× mais importante que o pior, "
        "o melhor deve ser cerca de 6× mais importante que o pior.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("8. Comentários (opcional)", 1)
    _table(doc, ["Observações"], [[""], [""], [""]], [17])

    p = doc.add_paragraph("Transcrição: a linha BO e a linha OW deste questionário correspondem às duas linhas do "
                          "decisor na folha Comparacoes do modelo Excel.")
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
