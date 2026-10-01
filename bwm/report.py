"""Relatórios: HTML interativo (Plotly), Word (.docx) e Excel com todas as tabelas."""

from __future__ import annotations

import html
import io
from datetime import datetime

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from . import analysis as A
from . import charts as C
from . import ranking as RK

REFERENCES = [
    "Rezaei, J. (2015). Best-worst multi-criteria decision-making method. Omega, 53, 49-57.",
    "Rezaei, J. (2016). Best-worst multi-criteria decision-making method: Some properties and a linear model. Omega, 64, 126-130.",
    "Mohammadi, M., & Rezaei, J. (2020). Bayesian best-worst method: A probabilistic group decision making model. Omega, 96, 102075.",
    "Liang, F., Brunelli, M., & Rezaei, J. (2020). Consistency issues in the best worst method: Measurements and thresholds. Omega, 96, 102175.",
    "Hwang, C. L., & Yoon, K. (1981). Multiple Attribute Decision Making: Methods and Applications. Springer (TOPSIS).",
    "Opricovic, S., & Tzeng, G. H. (2004). Compromise solution by MCDM methods: A comparative analysis of VIKOR and TOPSIS. European Journal of Operational Research, 156(2), 445-455.",
    "Hoffman, M. D., & Gelman, A. (2014). The No-U-Turn Sampler. Journal of Machine Learning Research, 15, 1593-1623.",
]

METHOD_TEXT = [
    "O Best-Worst Method (BWM) pede a cada decisor que identifique o critério mais importante (melhor) e o "
    "menos importante (pior) e que os compare com os restantes numa escala de 1 a 9. Obtêm-se dois vetores: "
    "Best-to-Others (quanto o melhor é preferido a cada critério) e Others-to-Worst (quanto cada critério é "
    "preferido ao pior). São precisas apenas 2n−3 comparações, contra n(n−1)/2 no AHP.",
    "O Bayesian BWM (Mohammadi & Rezaei, 2020) trata os vetores como dados de uma distribuição multinomial "
    "cujos parâmetros são os pesos de cada decisor. Os pesos individuais são ligados por uma distribuição "
    "Dirichlet centrada nos pesos agregados do grupo (w*), com concentração γ. Assim, o modelo estima "
    "simultaneamente os pesos individuais, o consenso do grupo e o grau de concordância, sem recorrer a médias "
    "que ignoram a incerteza.",
    "A estimação usa Monte Carlo via Cadeias de Markov (Hamiltonian Monte Carlo). Cada amostra é um conjunto "
    "plausível de pesos; a partir delas calculam-se médias, intervalos de credibilidade e o ranking credal: "
    "a probabilidade de um critério ser mais importante do que outro.",
    "A consistência dos julgamentos é avaliada com o rácio clássico CR = ξ*/CI(a_BW) (Rezaei, 2015), com o rácio "
    "baseado nos inputs (Liang et al., 2020) e com o valor ótimo ξ do modelo linear do BWM (Rezaei, 2016).",
    "As alternativas são ordenadas com TOPSIS, VIKOR e SAW. Cada método é aplicado a todas as amostras a posteriori "
    "dos pesos (simulação de Monte Carlo), o que dá a probabilidade de cada alternativa ficar em cada posição e "
    "permite ver se a recomendação é robusta à incerteza dos pesos e à escolha do método.",
]


def _fmt_df(df: pd.DataFrame, digits: int = 3) -> pd.DataFrame:
    out = df.copy()
    for c in out.columns:
        if pd.api.types.is_float_dtype(out[c]):
            out[c] = out[c].map(lambda v: "" if pd.isna(v) else f"{v:.{digits}f}")
    return out


def build_bundle(res, bo, ow, meta, alt=None, types=None, method="TOPSIS", v=0.5):
    """Calcula uma vez tudo o que os três formatos de relatório precisam."""
    cons = A.consistency_table(bo, ow)
    alt_res = A.alternatives_analysis(res, alt, types or {}, v) if alt is not None and len(alt) else None
    return {
        "meta": meta,
        "res": res,
        "bo": bo,
        "ow": ow,
        "weights": A.weights_table(res),
        "credal": A.credal_table(res),
        "chain": A.credal_chain(res),
        "cons": cons,
        "classic": A.classic_weights_table(bo, ow),
        "individual": A.individual_table(res),
        "distance": A.dm_distance(res),
        "diag": A.diagnostics_table(res),
        "alt": alt_res,
        "alt_raw": alt,
        "text": A.interpretation(res, cons, alt_res, method),
        "method": method,
        "types": types or {},
    }


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------

_CSS = """
:root{--pine:#2E5E4E;--ink:#1B2A26;--mist:#EEF2F0;--amber:#C98A1B}
*{box-sizing:border-box}
body{font-family:'Source Sans Pro','Segoe UI',system-ui,sans-serif;color:var(--ink);margin:0;background:#fff;line-height:1.55}
main{max-width:980px;margin:0 auto;padding:2.5rem 1.25rem 4rem}
header{border-left:6px solid var(--pine);padding:.4rem 0 .4rem 1.1rem;margin-bottom:2rem}
h1{font-size:2rem;margin:.2rem 0}
h2{font-size:1.35rem;margin-top:2.6rem;padding-bottom:.3rem;border-bottom:2px solid var(--mist)}
h3{font-size:1.05rem;margin-top:1.6rem}
p{max-width:72ch}
.meta{color:#4B5B56;font-size:.95rem}
.box{background:var(--mist);border-radius:6px;padding:1rem 1.2rem;margin:1rem 0}
.box li{margin:.35rem 0}
.tbl{overflow-x:auto;margin:.8rem 0}
table{border-collapse:collapse;font-size:.88rem;min-width:60%}
th,td{padding:.4rem .6rem;border-bottom:1px solid #DDE5E1;text-align:right;white-space:nowrap}
th:first-child,td:first-child{text-align:left}
th{background:var(--pine);color:#fff;font-weight:600}
.cap{font-size:.85rem;color:#4B5B56;margin-top:-.3rem}
.chain{font-size:1.05rem}
.chain b{color:var(--pine)}
.chain span{color:var(--amber);font-size:.85rem}
footer{margin-top:3rem;font-size:.85rem;color:#4B5B56}
@media print{h2{page-break-after:avoid}.plot{page-break-inside:avoid}}
"""


def _tbl(df, digits=3, index=False):
    return '<div class="tbl">' + _fmt_df(df, digits).to_html(index=index, border=0, escape=True) + "</div>"


def _fig(fig, first, offline=True):
    mode = (True if offline else "cdn") if first else False
    return '<div class="plot">' + fig.to_html(full_html=False, include_plotlyjs=mode,
                                               config={"displaylogo": False}) + "</div>"


def html_report(b, offline: bool = True) -> bytes:
    """offline=True embute a biblioteca Plotly (ficheiro maior, funciona sem Internet)."""
    res, meta = b["res"], b["meta"]
    s = res.summary()
    e = html.escape
    parts = [f"<!doctype html><html lang='pt-PT'><head><meta charset='utf-8'>"
             f"<meta name='viewport' content='width=device-width,initial-scale=1'>"
             f"<title>Relatório Bayesian BWM — {e(meta['title'])}</title><style>{_CSS}</style></head><body><main>"]
    parts.append(f"<header><div class='meta'>Relatório Bayesian Best-Worst Method</div>"
                 f"<h1>{e(meta['title'])}</h1><div class='meta'>Setor: {e(meta['industry'])} &nbsp;|&nbsp; "
                 f"Gerado em {datetime.now():%d/%m/%Y %H:%M}</div></header>")
    if meta.get("context"):
        parts.append(f"<p>{e(meta['context'])}</p>")

    parts.append("<h2>Síntese dos resultados</h2><div class='box'><ul>"
                 + "".join(f"<li>{e(t)}</li>" for t in b["text"]) + "</ul></div>")
    parts.append("<p class='chain'>Ranking credal: " + " ".join(
        f"<b>{e(a)}</b> <span>›{p:.2f}›</span>" for a, _, p in b["chain"])
        + f" <b>{e(b['chain'][-1][1])}</b></p>" if b["chain"] else "")
    parts.append("<p class='cap'>Os números entre os critérios são a probabilidade de o critério à esquerda "
                 "ser mais importante do que o seguinte.</p>")

    parts.append("<h2>1. Dados de entrada</h2>")
    parts.append(f"<p>{len(res.criteria)} critérios e {len(res.decision_makers)} decisores.</p>")
    crit_df = meta.get("criteria_df")
    if crit_df is not None:
        parts.append("<h3>Critérios</h3>" + _tbl(crit_df))
    parts.append("<h3>Vetores Best-to-Others</h3>" + _tbl(b["bo"], 0, index=True))
    parts.append("<h3>Vetores Others-to-Worst</h3>" + _tbl(b["ow"], 0, index=True))

    parts.append("<h2>2. Consistência dos julgamentos</h2>" + _tbl(b["cons"]))
    parts.append("<p class='cap'>CR input-based ≤ limiar indica consistência aceitável. "
                 "ξ próximo de 0 indica julgamentos coerentes no modelo linear do BWM.</p>")

    parts.append("<h2>3. Pesos agregados do grupo</h2>" + _tbl(b["weights"]))
    first = True
    parts.append(_fig(C.weights_bar(res.criteria, s), first, offline))
    first = False
    parts.append(_fig(C.posterior_violin(res.criteria, res.w_star), first))

    parts.append("<h2>4. Ranking credal</h2><p>Cada célula mostra a probabilidade de o critério da linha ser "
                 "mais importante do que o da coluna. Valores perto de 0,5 indicam empate prático.</p>")
    parts.append(_fig(C.credal_heatmap(res.criteria, res.credal_matrix()), first))
    parts.append(_fig(C.rank_prob_heatmap(res.criteria, res.rank_probabilities()), first))

    parts.append("<h2>5. Perspetiva de cada decisor</h2>")
    parts.append(_fig(C.individual_heatmap(res.criteria, res.decision_makers, res.individual_means(),
                                           res.w_star.mean(axis=0)), first))
    parts.append(_tbl(b["distance"]))
    parts.append("<p class='cap'>A distância de Jensen-Shannon (0 = igual ao consenso, 1 = totalmente diferente) "
                 "mostra que decisores se afastam mais da posição do grupo.</p>")
    parts.append(_fig(C.gamma_hist(res.gamma), first))
    parts.append("<h3>Comparação com o BWM clássico (modelo linear, por decisor)</h3>" + _tbl(b["classic"], 3, index=True))

    if b["alt"] is not None:
        ar, mth = b["alt"], b["method"]
        md = ar["methods"][mth]
        parts.append("<h2>6. Avaliação das alternativas</h2>")
        parts.append(f"<p>Método principal: <b>{mth}</b>. {e(RK.METHOD_INFO[mth])} Cada método é aplicado a "
                     "todas as amostras dos pesos, o que propaga a incerteza para o ranking final.</p>")
        parts.append("<h3>Desempenho original</h3>" + _tbl(b["alt_raw"], 2, index=True))
        parts.append(f"<h3>Resultado — {mth}</h3>" + _tbl(md["table"]))
        parts.append(_fig(C.alternatives_bar(ar["names"], md["raw"], md["metric"], mth == "VIKOR", mth), first))
        parts.append("<h3>Probabilidade de cada posição</h3>" + _tbl(md["rank_prob"], 2, index=True))
        parts.append("<h3>Comparação SAW · TOPSIS · VIKOR</h3>" + _tbl(ar["comparison"], 2))
        parts.append(_fig(C.method_first_prob(ar), first))
        parts.append("<p class='cap'>Correlação de Spearman entre os rankings dos métodos (1 = rankings iguais):</p>"
                     + _tbl(RK.spearman_matrix(ar), 2, index=True))
        vc = ar["vikor_check"]
        parts.append(f"<p><b>Condições do VIKOR</b> (pesos médios, v = {ar['v']:.2f}): vantagem aceitável "
                     f"{'cumprida' if vc['C1_vantagem'] else 'não cumprida'}; estabilidade "
                     f"{'cumprida' if vc['C2_estabilidade'] else 'não cumprida'}. {e(vc['mensagem'])}</p>")
        top = int(np.argmax(res.w_star.mean(axis=0)))
        g, raw, sw = RK.sensitivity(mth, ar["X"], ar["is_benefit"], res.w_star.mean(axis=0), top, ar["v"])
        parts.append("<h3>Análise de sensibilidade</h3><p>Variação do peso do critério mais importante, mantendo "
                     "as proporções entre os restantes. As linhas verticais pontilhadas marcam mudanças de líder.</p>")
        parts.append(_fig(C.sensitivity_chart(g, raw, ar["names"], res.criteria[top],
                                              float(res.w_star.mean(axis=0)[top]), md["metric"], sw), first))

    parts.append("<h2>Diagnóstico da estimação</h2>" + _tbl(b["diag"]))
    parts.append(f"<p class='cap'>Cadeias: {res.chains}; amostras por cadeia: {res.draws}; taxa de aceitação média: "
                 f"{res.acceptance.get('taxa_aceitacao', float('nan')):.2f}; divergências: "
                 f"{res.acceptance.get('divergencias', 0)}. R-hat &lt; 1,01 e ESS &gt; 400 indicam boa convergência.</p>")

    parts.append("<h2>Metodologia</h2>" + "".join(f"<p>{e(t)}</p>" for t in METHOD_TEXT))
    parts.append("<h2>Referências</h2><ul>" + "".join(f"<li>{e(r)}</li>" for r in REFERENCES) + "</ul>")
    parts.append("<footer>Gerado pela aplicação Bayesian BWM Studio.</footer></main></body></html>")
    return "".join(parts).encode("utf-8")


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------


def _doc_table(doc, df, digits=3, index=False):
    if index:
        df = df.copy()
        df.index.name = df.index.name or ""
        df = df.reset_index()
    df = _fmt_df(df, digits)
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(df.columns):
        t.rows[0].cells[i].text = str(c)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)
    doc.add_paragraph()


def docx_report(b) -> bytes:
    res, meta = b["res"], b["meta"]
    s = res.summary()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2)

    h = doc.add_heading(meta["title"], 0)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph(f"Relatório Bayesian Best-Worst Method  |  Setor: {meta['industry']}  |  "
                          f"{datetime.now():%d/%m/%Y}")
    p.runs[0].font.color.rgb = RGBColor(0x4B, 0x5B, 0x56)
    if meta.get("context"):
        doc.add_paragraph(meta["context"])

    doc.add_heading("Síntese dos resultados", 1)
    for t in b["text"]:
        doc.add_paragraph(t, style="List Bullet")
    if b["chain"]:
        doc.add_paragraph("Ranking credal: " + "  ›  ".join(
            f"{a} ({p:.2f})" for a, _, p in b["chain"]) + f"  ›  {b['chain'][-1][1]}")

    doc.add_heading("1. Dados de entrada", 1)
    if meta.get("criteria_df") is not None:
        doc.add_heading("Critérios", 2)
        _doc_table(doc, meta["criteria_df"])
    doc.add_heading("Vetores Best-to-Others", 2)
    _doc_table(doc, b["bo"], 0, index=True)
    doc.add_heading("Vetores Others-to-Worst", 2)
    _doc_table(doc, b["ow"], 0, index=True)

    doc.add_heading("2. Consistência dos julgamentos", 1)
    _doc_table(doc, b["cons"])

    doc.add_heading("3. Pesos agregados do grupo", 1)
    _doc_table(doc, b["weights"])
    doc.add_picture(io.BytesIO(C.mpl_weights(res.criteria, s)), width=Cm(15))

    doc.add_heading("4. Ranking credal", 1)
    doc.add_paragraph("Probabilidade de o critério da linha ser mais importante do que o da coluna.")
    doc.add_picture(io.BytesIO(C.mpl_heatmap(res.credal_matrix(), res.criteria, res.criteria,
                                             "P(linha > coluna)")), width=Cm(15))

    doc.add_heading("5. Perspetiva de cada decisor", 1)
    _doc_table(doc, b["individual"], 3, index=True)
    _doc_table(doc, b["distance"])
    doc.add_heading("BWM clássico (modelo linear)", 2)
    _doc_table(doc, b["classic"], 3, index=True)

    if b["alt"] is not None:
        ar, mth = b["alt"], b["method"]
        md = ar["methods"][mth]
        doc.add_heading("6. Avaliação das alternativas", 1)
        doc.add_paragraph(f"Método principal: {mth}. {RK.METHOD_INFO[mth]}")
        _doc_table(doc, b["alt_raw"], 2, index=True)
        doc.add_heading(f"Resultado — {mth}", 2)
        _doc_table(doc, md["table"])
        doc.add_picture(io.BytesIO(C.mpl_alternatives(ar["names"], md["raw"], md["metric"], mth == "VIKOR")),
                        width=Cm(15))
        doc.add_heading("Comparação SAW, TOPSIS e VIKOR", 2)
        _doc_table(doc, ar["comparison"], 2)
        doc.add_picture(io.BytesIO(C.mpl_method_first(ar)), width=Cm(15))
        doc.add_paragraph("Condições do VIKOR: " + ar["vikor_check"]["mensagem"])

    doc.add_heading("Diagnóstico da estimação", 1)
    _doc_table(doc, b["diag"])
    doc.add_paragraph(f"Cadeias: {res.chains}; amostras por cadeia: {res.draws}; divergências: "
                      f"{res.acceptance.get('divergencias', 0)}.")

    doc.add_heading("Metodologia", 1)
    for t in METHOD_TEXT:
        doc.add_paragraph(t)
    doc.add_heading("Referências", 1)
    for r in REFERENCES:
        doc.add_paragraph(r, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def excel_report(b) -> bytes:
    buf = io.BytesIO()
    res = b["res"]
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame({"Síntese": b["text"]}).to_excel(xw, sheet_name="Sintese", index=False)
        b["weights"].to_excel(xw, sheet_name="Pesos agregados", index=False)
        b["credal"].to_excel(xw, sheet_name="Ranking credal")
        pd.DataFrame(res.rank_probabilities(), index=res.criteria,
                     columns=[f"{r + 1}.º" for r in range(len(res.criteria))]).to_excel(xw, sheet_name="Prob. posicao")
        b["individual"].to_excel(xw, sheet_name="Pesos individuais")
        b["distance"].to_excel(xw, sheet_name="Distancia consenso", index=False)
        b["cons"].to_excel(xw, sheet_name="Consistencia", index=False)
        b["classic"].to_excel(xw, sheet_name="BWM classico")
        b["bo"].to_excel(xw, sheet_name="Dados BO")
        b["ow"].to_excel(xw, sheet_name="Dados OW")
        if b["alt"] is not None:
            b["alt"]["comparison"].to_excel(xw, sheet_name="Alt comparacao metodos", index=False)
            for mth, d in b["alt"]["methods"].items():
                d["table"].to_excel(xw, sheet_name=f"Alt {mth}", index=False)
                d["rank_prob"].to_excel(xw, sheet_name=f"Alt {mth} prob posicao")
        b["diag"].to_excel(xw, sheet_name="Diagnostico", index=False)
        pd.DataFrame(res.w_star[:: max(1, len(res.w_star) // 2000)], columns=res.criteria) \
            .to_excel(xw, sheet_name="Amostras w_estrela", index=False)
    return buf.getvalue()


def samples_csv(res) -> bytes:
    df = pd.DataFrame(res.w_star, columns=res.criteria)
    df["gamma"] = res.gamma
    df.insert(0, "cadeia", np.repeat(np.arange(1, res.chains + 1), res.draws))
    return df.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig")
